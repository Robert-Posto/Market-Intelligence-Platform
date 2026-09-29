"""Încarcă catalogul intern de produse Libra (Excel) în tabela `catalog_libra`.

De ce: produsele Libra nu se mai colectează de pe web. Catalogul intern
(export din Sales Command Center) e REFERINȚA după care se va colecta ulterior
de la celelalte bănci. Rândurile nu intră în `observations`: acolo stau doar
valorile observate pe web, cu citat și sursă.

Fidelitate: fiecare rând din foaia „Catalog produse" ajunge în bază cu
coloanele structurate (goluri -> NULL) ȘI cu rândul complet în `rand_brut`
(jsonb, numele coloanelor exact ca în antet). Nu se interpretează nimic peste
ce scrie în fișier, în afară de două derivări simple și reversibile:
`segment` (din „Adresabilitate") și `categorie_cod` (codul din paranteza
categoriei); ambele lasă textul original alături.

Idempotent: cheia e (fișier, foaie, rând). Reîncărcarea aceluiași fișier
actualizează rândurile existente și șterge doar rândurile acelui fișier care
nu mai există în el; alte fișiere (alte exporturi) rămân neatinse.

Rulare:
    python ingest/load_catalog_libra.py                 # fișierul implicit
    python ingest/load_catalog_libra.py cale/fisier.xlsx
    python ingest/load_catalog_libra.py --uscat         # doar afișează
"""

import argparse
import collections
import datetime
import json
import os
import re
import sys

import psycopg2
import psycopg2.extras

AICI = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.dirname(AICI), AICI]
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import config                      # noqa: E402
import normalizeaza as N           # noqa: E402

FISIER_IMPLICIT = r"C:\Users\robert.postolache\Downloads\Catalog produse 2026-09-28.xlsx"
FOAIE_CATALOG = "Catalog produse"
FOAIE_SUMAR = "Sumar"
BANCA = "libra"

# antet din Excel -> coloana din tabelă (ce nu e aici rămâne doar în rand_brut)
COLOANE = {
    "Cod": "cod",
    "Denumire": "denumire",
    "Categorie produs": "categorie",
    "Tip produs": "tip_produs",
    "Adresabilitate": "adresabilitate",
    "Caracteristicile produsului": "caracteristici",
    "Criterii de eligibilitate": "criterii_eligibilitate",
    "Când să recomand acest produs": "cand_recomand",
    "Beneficiu pentru client": "beneficiu_client",
    "Sursa / document intern": "sursa_document",
    "Linii de business eligibile": "linii_business",
    "CAEN eligibile (prefixe)": "caen_prefixe",
}
NUMERICE = {
    "Cifră de afaceri minimă (lei)": "cifra_afaceri_min",
    "Cifră de afaceri maximă (lei)": "cifra_afaceri_max",
    "Angajați minim": "angajati_min",
    "Vechime minimă firmă (ani)": "vechime_min_ani",
    "Vechime maximă firmă (ani)": "vechime_max_ani",
}
SEGMENT = {
    "doar persoane fizice": "PF",
    "doar persoane juridice": "PJ",
    "oricine (pf și pj)": "PF+PJ",
}
RE_COD_CATEGORIE = re.compile(r"\(([A-Z0-9_]+)\)\s*$")

COL_INSERT = (["id_banca", "produs_prioritar", "segment", "categorie_cod", "rand_brut",
               "fisier", "foaie", "rand", "generat_la"]
              + list(COLOANE.values()) + list(NUMERICE.values()))


def _text(v):
    """Text tăiat la capete; gol -> None. Restul valorilor -> str."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _numar(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str):
        v = v.strip().replace(",", ".")
    n = float(v)
    return int(n) if n == int(n) else n


def _json_sigur(v):
    """Valoarea brută, serializabilă în jsonb (date -> ISO)."""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


def genereaza_la(ws_sumar):
    """„Generat la" din foaia Sumar („28.09.2026, 14:29:51") -> datetime, sau None."""
    if ws_sumar is None:
        return None
    for eticheta, valoare in ws_sumar.iter_rows(values_only=True):
        if _text(eticheta) == "Generat la":
            if isinstance(valoare, datetime.datetime):
                return valoare
            try:
                return datetime.datetime.strptime(_text(valoare), "%d.%m.%Y, %H:%M:%S")
            except (TypeError, ValueError):
                return None
    return None


def parseaza_rand(antet, valori, nr_rand, fisier, foaie, generat_la=None):
    """Un rând din foaie -> dict pentru tabelă. Fără bază de date, testabil."""
    brut = {h: _json_sigur(v) for h, v in zip(antet, valori) if h}
    r = {"fisier": fisier, "foaie": foaie, "rand": nr_rand, "generat_la": generat_la,
         "rand_brut": brut}
    for h, col in COLOANE.items():
        r[col] = _text(brut.get(h))
    for h, col in NUMERICE.items():
        r[col] = _numar(brut.get(h))
    r["produs_prioritar"] = (_text(brut.get("Produs prioritar")) or "").lower() == "da"
    r["segment"] = SEGMENT.get((r["adresabilitate"] or "").lower())
    m = RE_COD_CATEGORIE.search(r["categorie"] or "")
    r["categorie_cod"] = m.group(1) if m else None
    return r


def citeste(cale):
    """Rândurile foii „Catalog produse"; rândurile complet goale se sar."""
    import openpyxl
    wb = openpyxl.load_workbook(cale, data_only=True, read_only=True)
    if FOAIE_CATALOG not in wb.sheetnames:
        raise SystemExit(f"Lipsește foaia „{FOAIE_CATALOG}” (foi: {wb.sheetnames})")
    generat = genereaza_la(wb[FOAIE_SUMAR]) if FOAIE_SUMAR in wb.sheetnames else None
    fisier = os.path.basename(cale)
    randuri = []
    for nr, valori in enumerate(wb[FOAIE_CATALOG].iter_rows(values_only=True), 1):
        if nr == 1:
            antet = [_text(h) for h in valori]
            continue
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in valori):
            continue
        r = parseaza_rand(antet, valori, nr, fisier, FOAIE_CATALOG, generat)
        if not r["cod"] or not r["denumire"]:
            raise SystemExit(f"Rândul {nr}: fără Cod sau Denumire — nu inventez, verifică fișierul")
        randuri.append(r)
    wb.close()          # read_only ține fișierul deschis (pe Windows nu se mai poate șterge)
    coduri = collections.Counter(r["cod"] for r in randuri)
    dubluri = [c for c, n in coduri.items() if n > 1]
    if dubluri:
        print(f"ATENȚIE: coduri repetate în fișier: {dubluri}", file=sys.stderr)
    return randuri


def _json_ro(o):
    return json.dumps(o, ensure_ascii=False)


def scrie(randuri):
    """Upsert pe (fisier, foaie, rand) + ștergerea rândurilor dispărute din fișier."""
    if not randuri:
        raise SystemExit("Niciun rând de încărcat.")
    fisier, foaie = randuri[0]["fisier"], randuri[0]["foaie"]
    with psycopg2.connect(N.dsn()) as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM banci WHERE slug = %s", (BANCA,))
            id_banca = cur.fetchone()[0]
            valori = []
            for r in randuri:
                r = dict(r, id_banca=id_banca,
                         rand_brut=psycopg2.extras.Json(r["rand_brut"], dumps=_json_ro))
                valori.append(tuple(r[c] for c in COL_INSERT))
            actualizeaza = ", ".join(f"{c} = EXCLUDED.{c}" for c in COL_INSERT
                                     if c not in ("fisier", "foaie", "rand")) + ", importat_la = now()"
            psycopg2.extras.execute_values(
                cur,
                f"INSERT INTO catalog_libra ({', '.join(COL_INSERT)}) VALUES %s "
                f"ON CONFLICT (fisier, foaie, rand) DO UPDATE SET {actualizeaza}",
                valori)
            cur.execute("DELETE FROM catalog_libra WHERE fisier = %s AND foaie = %s "
                        "AND NOT (rand = ANY(%s))",
                        (fisier, foaie, [r["rand"] for r in randuri]))
            return cur.rowcount


def rezumat(randuri):
    print(f"{len(randuri)} produse din „{randuri[0]['fisier']}” / „{randuri[0]['foaie']}”")
    for titlu, cheie in (("categorie", "categorie_cod"), ("segment", "segment")):
        print(f"\nPe {titlu}:")
        for k, n in collections.Counter(r[cheie] or "(gol)" for r in randuri).most_common():
            print(f"  {n:4d}  {k}")
    print(f"\nProduse prioritare: {sum(r['produs_prioritar'] for r in randuri)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("fisier", nargs="?", default=FISIER_IMPLICIT, help="calea fișierului Excel")
    ap.add_argument("--uscat", action="store_true",
                    help="doar citește și afișează; nu scrie nimic în bază")
    a = ap.parse_args()

    config.incarca()
    randuri = citeste(a.fisier)
    rezumat(randuri)
    if a.uscat:
        print("\n--uscat: nu s-a scris nimic.")
        return 0
    sterse = scrie(randuri)
    print(f"\nScrise {len(randuri)} rânduri în catalog_libra ({sterse} dispărute din fișier, șterse).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
